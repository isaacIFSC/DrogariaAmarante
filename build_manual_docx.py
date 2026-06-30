from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\isaac\OneDrive\Documentos\Projetos\DrogariaAmarante")
OUTPUT_DIR = ROOT / "output" / "docx"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT = OUTPUT_DIR / "Manual_do_Usuario_Drogaria_Amarante.docx"


TITLE = "Manual de acesso à plataforma Drogaria Amarante"
SUBTITLE = "Guia do usuário"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_paragraph_border_bottom(paragraph, color="D9E2F0", size="8", space="1"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_char_begin)

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    run._r.append(instr_text)

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_end)


def set_run_font(run, size=11, bold=False, italic=False, color="000000", name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc, text="", *, style=None, size=11, bold=False, italic=False, color="000000",
                  align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, line=1.25, name="Calibri"):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color, name=name)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    if level == 1:
        fmt.space_before = Pt(18)
        fmt.space_after = Pt(10)
    elif level == 2:
        fmt.space_before = Pt(12)
        fmt.space_after = Pt(6)
    else:
        fmt.space_before = Pt(8)
        fmt.space_after = Pt(4)
    fmt.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size={1: 16, 2: 13}.get(level, 12), bold=False, color="000000")
    return p


def add_figure_caption(doc, text):
    return add_paragraph(
        doc,
        text,
        size=9,
        italic=True,
        color="000000",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=2,
        after=10,
        line=1.0,
    )


def add_note_box(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    set_cell_width(cell, 6.5)
    set_cell_shading(cell, "F4F7FB")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=10, color="000000")
    return table


def set_footers(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    stop = Inches(6.5)
    p.paragraph_format.tab_stops.add_tab_stop(stop, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)
    r1 = p.add_run("Drogaria Amarante")
    set_run_font(r1, size=9, color="000000")
    p.add_run("\t")
    add_page_field(p)
    for run in p.runs:
        if run._r.find(qn("w:fldChar")) is not None or run._r.find(qn("w:instrText")) is not None:
            run.font.size = Pt(9)


def add_section_break(doc):
    doc.add_page_break()


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    for name, size, color in [("Heading 1", 16, "000000"), ("Heading 2", 13, "000000"), ("Heading 3", 12, "000000")]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)

    # Cover page
    first = doc.add_paragraph()
    first.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first.paragraph_format.space_before = Pt(180)
    first.paragraph_format.space_after = Pt(6)
    first_run = first.add_run(TITLE)
    set_run_font(first_run, size=24, bold=False, color="000000")

    subtitle = add_paragraph(
        doc,
        SUBTITLE,
        size=13,
        italic=False,
        color="000000",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=0,
        after=16,
        line=1.0,
    )

    meta = add_paragraph(
        doc,
        "Documento de orientação para acesso e uso da plataforma",
        size=10.5,
        color="000000",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=0,
        after=0,
        line=1.0,
    )

    add_note_box(
        doc,
        "Versão de entrega acadêmica com estrutura inspirada em manual de acesso: sumário, figuras, instruções de uso e área administrativa."
    )

    add_section_break(doc)

    # Registro de revisoes
    add_heading(doc, "Registro de revisões", 1)
    add_paragraph(
        doc,
        "Data | Responsável | Descrição",
        size=10.5,
        bold=True,
        color="4F5B66",
        after=4,
        line=1.0,
    )

    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.2, 1.4, 3.9]
    headers = ["Data", "Responsável", "Descrição"]
    row = table.rows[0]
    for i, h in enumerate(headers):
        cell = row.cells[i]
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, "E8EEF5")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h)
        set_run_font(r, size=10.5, bold=True, color="000000")

    row = table.rows[1]
    values = ["28/06/2026", "Projeto", "Versão inicial do manual de usuário baseada no modelo do professor."]
    for i, v in enumerate(values):
        cell = row.cells[i]
        set_cell_width(cell, widths[i])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 2 else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(v)
        set_run_font(r, size=10.5, color="000000")

    add_section_break(doc)

    # Sumario
    add_heading(doc, "Sumário", 1)
    summary_lines = [
        "1. Introdução ............................................................................................ 4",
        "2. Página inicial da plataforma .............................................................. 4",
        "2.1. Barra superior e menu ................................................................. 4",
        "2.2. Listagem de produtos ............................................................... 4",
        "3. Primeiro acesso ............................................................................... 4",
        "3.1. Criando conta ............................................................................. 4",
        "3.2. Fazendo login .............................................................................. 5",
        "3.3. Acessando o perfil ....................................................................... 5",
        "4. Área administrativa ........................................................................ 5",
        "4.1. Cadastro manual de produtos .................................................... 5",
        "4.2. Importação de produtos por CSV ................................................. 5",
        "5. Página de produto ........................................................................... 6",
        "6. Carrinho de compras ..................................................................... 6",
        "7. Observações importantes .............................................................. 6",
    ]
    for line in summary_lines:
        add_paragraph(doc, line, size=10.5, color="000000", after=2, line=1.05)

    add_heading(doc, "Figuras", 1)
    fig_lines = [
        "Figura 1: Página inicial da Drogaria Amarante",
        "Figura 2: Menu principal da plataforma",
        "Figura 3: Cards de produtos na pagina inicial",
        "Figura 4: Tela de cadastro de usuário",
        "Figura 5: Tela de login",
        "Figura 6: Página de perfil do usuário",
        "Figura 7: Área administrativa com formulário de cadastro",
        "Figura 8: Importação de produtos por CSV",
        "Figura 9: Página de detalhes do produto",
        "Figura 10: Página do carrinho de compras",
        "Figura 11: Estrutura geral da plataforma",
    ]
    for line in fig_lines:
        add_paragraph(doc, line, size=10, color="555555", after=1, line=1.05)

    add_section_break(doc)

    # Body
    add_heading(doc, "1. Introdução", 1)
    add_paragraph(
        doc,
        "Este manual foi preparado para orientar o usuário no uso da plataforma Drogaria Amarante. "
        "Aqui você encontra as instruções para criar conta, entrar no sistema, consultar produtos, "
        "visualizar o perfil e acessar a área administrativa de cadastro de produtos.",
        after=8,
    )

    add_heading(doc, "2. Página inicial da plataforma", 1)
    add_paragraph(
        doc,
        "A página inicial da Drogaria Amarante pode ser acessada pelo endereço raiz da aplicação, ou seja, pelo caminho /.",
        after=6,
    )
    add_paragraph(
        doc,
        "Nela você encontra o banner promocional, o campo de busca, os atalhos de acesso e a lista de produtos carregada automaticamente pela API.",
        after=8,
    )
    add_figure_caption(doc, "Figura 1: Página inicial da Drogaria Amarante")

    add_heading(doc, "2.1. Barra superior e menu", 2)
    add_paragraph(
        doc,
        "Na parte superior da página há uma faixa promocional e um menu com atalhos rápidos. "
        "Os itens mais importantes são Entrar, Meu perfil, Cadastrar produto, Sair e Carrinho.",
        after=8,
    )
    add_figure_caption(doc, "Figura 2: Menu principal da plataforma")

    add_heading(doc, "2.2. Listagem de produtos", 2)
    add_paragraph(
        doc,
        "Os produtos exibidos na página inicial são carregados automaticamente pela API /api/produtos.",
        after=4,
    )
    add_paragraph(
        doc,
        "Cada card apresenta imagem, categoria, marca, nome, descrição e preço. Ao clicar em um card, o usuário acessa a página de detalhes do produto.",
        after=8,
    )
    add_figure_caption(doc, "Figura 3: Cards de produtos na página inicial")

    add_heading(doc, "3. Primeiro acesso", 1)
    add_paragraph(
        doc,
        "Para utilizar a plataforma pela primeira vez, o usuário deve criar uma conta e depois fazer login.",
        after=8,
    )

    add_heading(doc, "3.1. Criando conta", 2)
    add_paragraph(
        doc,
        "Para criar uma conta, acesse o endereço /cadastro. Preencha nome, sobrenome, e-mail, senha, CPF, telefone, gênero e data de nascimento. "
        "Depois clique em Cadastrar. Ao final do processo, o sistema redireciona para a tela de login.",
        after=8,
    )
    add_figure_caption(doc, "Figura 4: Tela de cadastro de usuário")

    add_heading(doc, "3.2. Fazendo login", 2)
    add_paragraph(
        doc,
        "Para entrar na plataforma, acesse /login. No formulário de login, informe E-mail ou CPF e Senha, e depois clique em Entrar. "
        "Se os dados estiverem corretos, o sistema abre a página inicial ou a página solicitada anteriormente.",
        after=8,
    )
    add_paragraph(
        doc,
        "Observação: o sistema contém um usuário padrão criado automaticamente para uso administrativo. "
        "E-mail: admin@drugstoreamarante.com. Senha: 123456.",
        after=8,
    )
    add_figure_caption(doc, "Figura 5: Tela de login")

    add_heading(doc, "3.3. Acessando o perfil", 2)
    add_paragraph(
        doc,
        "Depois de autenticado, o usuário pode visualizar seus dados em /perfil. Nessa página são exibidos nome, sobrenome, e-mail, CPF, telefone e data de nascimento. "
        "Também há atalhos para voltar à página inicial, sair da conta e acessar a área administrativa.",
        after=8,
    )
    add_figure_caption(doc, "Figura 6: Página de perfil do usuário")

    add_heading(doc, "4. Área administrativa", 1)
    add_paragraph(
        doc,
        "A área administrativa fica disponível no endereço /admin/produtos. Para acessar essa área, o usuário precisa estar logado. Caso não esteja, o sistema envia para a tela de login com retorno automático para a página administrativa.",
        after=8,
    )

    add_heading(doc, "4.1. Cadastro manual de produtos", 2)
    add_paragraph(
        doc,
        "O formulário de cadastro manual permite incluir um produto por vez. Os campos obrigatórios são nome, marca, categoria, preço, estoque e descrição. "
        "O campo de imagem é opcional. Se ele não for preenchido, o sistema utiliza uma imagem padrão. Depois de preencher o formulário, clique em Salvar produto.",
        after=8,
    )
    add_figure_caption(doc, "Figura 7: Área administrativa com formulário de cadastro")

    add_heading(doc, "4.2. Importação de produtos por CSV", 2)
    add_paragraph(
        doc,
        "A mesma área administrativa permite importar vários produtos por arquivo CSV. O sistema aceita dois formatos: cadastro padrão, com colunas como nome, descrição, preço, estoque, categoria e marca; "
        "e relatório de estoque, com colunas como código, descrição, NCM, ICMS, unidade, qtd, vlrUnit e vlrTotal.",
        after=8,
    )
    add_paragraph(
        doc,
        "Para importar, selecione o arquivo CSV, clique em Importar arquivo e aguarde a mensagem de retorno na tela.",
        after=8,
    )
    add_figure_caption(doc, "Figura 8: Importação de produtos por CSV")

    add_heading(doc, "5. Página de produto", 1)
    add_paragraph(
        doc,
        "A página de produto pode ser acessada pelo caminho /produto/{id}. Ela mostra os detalhes completos do item selecionado: imagem, categoria, nome, marca, descrição, preço, estoque e status de disponibilidade.",
        after=8,
    )
    add_figure_caption(doc, "Figura 9: Página de detalhes do produto")

    add_heading(doc, "6. Carrinho de compras", 1)
    add_paragraph(
        doc,
        "A página /carrinho apresenta o carrinho de compras e o resumo do pedido. Nesta tela o usuário vê itens adicionados, preço unitário, quantidade, total por item, subtotal, descontos, taxa de entrega e total final.",
        after=8,
    )
    add_paragraph(
        doc,
        "Observação importante: no estado atual do projeto, o carrinho funciona como uma tela demonstrativa. O fluxo completo de adicionar, remover e finalizar compra ainda não está integrado.",
        after=8,
    )
    add_figure_caption(doc, "Figura 10: Página do carrinho de compras")

    add_heading(doc, "7. Observações importantes", 1)
    add_paragraph(
        doc,
        "A plataforma foi desenhada para uso em ambiente web. A busca e a listagem de produtos dependem da API de produtos. A área administrativa exige sessão logada.",
        after=6,
    )
    add_paragraph(
        doc,
        "Se ocorrer erro no login, o sistema retorna para a tela de entrada com aviso de acesso inválido. Se não houver produtos cadastrados, a página inicial exibe uma mensagem indicando que não há itens para mostrar.",
        after=8,
    )
    add_figure_caption(doc, "Figura 11: Estrutura geral da plataforma")

    set_footers(section)
    return doc


def main():
    doc = build_document()
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
